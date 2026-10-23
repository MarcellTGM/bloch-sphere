"""Egyszerű Markdown -> DOCX átalakító a Bloch-gömb dokumentációhoz.

Támogatott elemek: címek (#, ##, ###), bekezdések, felsorolások (-),
számozott listák (1.), táblázatok (| ... |), kódrészletek (```), képek
(![alt](utvonal)), kiemelés (**félkövér**, `kód`, *dőlt*), blokkidézet (>).
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)


def unescape(text):
    return text.replace("\\|", "|").replace("\\_", "_").replace("\\*", "*")


def add_inline(paragraph, text, bold_all=False, italic_all=False, size=None):
    """Inline formázás: **félkövér**, `kód`, *dőlt*."""
    text = unescape(text)
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")
    position = 0

    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold = bold_all
            run.italic = italic_all
            if size:
                run.font.size = Pt(size)

        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic_all
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.bold = bold_all
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.bold = bold_all

        if size:
            run.font.size = Pt(size)
        position = match.end()

    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = bold_all
        run.italic = italic_all
        if size:
            run.font.size = Pt(size)


def add_table(document, rows):
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, cell_text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_inline(cell.paragraphs[0], cell_text, bold_all=True, size=9)
        properties = cell._tc.get_or_add_tcPr()
        shading = properties.makeelement(qn("w:shd"), {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): "DCE6F1",
        })
        properties.append(shading)

    for line in body:
        cells = table.add_row().cells
        for index, cell_text in enumerate(line):
            if index >= len(cells):
                break
            cells[index].text = ""
            add_inline(cells[index].paragraphs[0], cell_text, size=9)

    document.add_paragraph()
    return table


def convert(markdown_path, docx_path, image_dir=None):
    base = image_dir or os.path.dirname(os.path.abspath(markdown_path))
    lines = open(markdown_path, encoding="utf-8-sig").read().splitlines()

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(4)

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            document.add_heading(unescape(text), level=min(level, 3))
            index += 1
            continue

        if stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.+?)\)", stripped)
            if match:
                alt, path = match.group(1), match.group(2)
                full = path if os.path.isabs(path) else os.path.join(base, path)
                if os.path.exists(full):
                    document.add_picture(full, width=Cm(15.5))
                    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption = document.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption.add_run(alt)
                    run.italic = True
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = MUTED
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while index < len(lines):
                probe = lines[index].strip()
                if not (probe.startswith("|") and probe.endswith("|")):
                    break
                if re.fullmatch(r"\|[\s:\-|]+\|", probe):
                    index += 1
                    continue
                rows.append([cell.strip() for cell in probe.strip("|").split("|")])
                index += 1
            if rows:
                add_table(document, rows)
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            add_inline(paragraph, stripped.lstrip("> ").strip(), italic_all=True, size=10)
            index += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    document.save(docx_path)
    print(f"KESZ: {docx_path}")


if __name__ == "__main__":
    source = sys.argv[1]
    target = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(source)[0] + ".docx"
    images = sys.argv[3] if len(sys.argv) > 3 else None
    convert(source, target, images)

